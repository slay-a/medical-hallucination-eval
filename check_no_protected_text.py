#!/usr/bin/env python3
"""
check_no_protected_text.py — Verify that no tracked repository file contains text from the credentialed sources.

Builds hashed 8-word shingles from every credentialed text used in this thesis (ann-pt-summ hospital courses and summaries,
MedNLI sentences) and from every MIMIC-derived output kept in results_private/ (main-study generations, excerpts, answers,
decompositions), then scans every file tracked by git, including the text of PDF, Word and PowerPoint files.
Shingles that also occur in the public MTSamples corpus are ignored, so that generic phrases shared with public text do not
count.  Matches are reported separately for the credentialed source texts and for the private model outputs, with the number
of credentialed documents that contain each matching shingle (a phrase found in many notes is generic wording, not a copied
passage).  Writes results/protected_text_scan.csv (per file: shingles, matches with the source texts, matches with the model
outputs only, largest source-document frequency) and results/protected_text_scan_summary.csv (totals over all files, with
distinct matching shingles counted once), and prints only counts.  Every file version reachable from any commit is scanned
as well (results/protected_text_scan_history.csv; skip with --no-history).
"""
import glob
import hashlib
import json
import re
import subprocess
import sys
from pathlib import Path

import pandas as pd

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
K = 8
WORD = re.compile(r"[a-z0-9]+")


def shingles(text):
    w = WORD.findall(str(text).lower())
    return {hashlib.blake2b(" ".join(w[i:i + K]).encode(), digest_size=8).digest() for i in range(max(0, len(w) - K + 1))}


def file_text(path):
    p = str(path)
    try:
        if p.endswith(".pdf"):
            return subprocess.run(["pdftotext", p, "-"], capture_output=True, text=True).stdout
        if p.endswith(".docx"):
            import docx
            return "\n".join(par.text for par in docx.Document(p).paragraphs)
        if p.endswith(".pptx"):
            from pptx import Presentation
            return "\n".join(sh.text_frame.text for sl in Presentation(p).slides for sh in sl.shapes if sh.has_text_frame)
        if p.endswith((".png", ".jpg", ".gz", ".zip")):
            return ""
        return Path(p).read_text(errors="ignore")
    except Exception:  # noqa: BLE001
        return ""


def history_texts():
    """(path, blob id, text) for every file version reachable from any commit."""
    import tempfile
    objs = subprocess.run(["git", "rev-list", "--all", "--objects"], cwd=HERE, capture_output=True, text=True).stdout.split("\n")
    seen = set()
    with tempfile.TemporaryDirectory() as tmp:
        for line in objs:
            parts = line.split(" ", 1)
            if len(parts) != 2 or parts[0] in seen:
                continue
            sha, path = parts; seen.add(sha)
            kind = subprocess.run(["git", "cat-file", "-t", sha], cwd=HERE, capture_output=True, text=True).stdout.strip()
            if kind != "blob":
                continue
            data = subprocess.run(["git", "cat-file", "-p", sha], cwd=HERE, capture_output=True).stdout
            f = Path(tmp) / (sha + Path(path).suffix)
            f.write_bytes(data)
            yield path, sha, file_text(f)
            f.unlink()


def main():
    from collections import Counter
    from calibration_annptsumm import DEFAULT_DATA, read_jsonl
    src = Counter()          # shingle -> number of credentialed source documents that contain it
    for f in ("hallucinations_mimic_di.jsonl", "hallucinations_mimic_di_validation.jsonl", "hallucinations_generated_di.jsonl"):
        for r in read_jsonl(Path(DEFAULT_DATA) / f):
            src.update(shingles(r["text"]) | shingles(r["summary"]))
    for f in glob.glob(str(Path.home() / "Desktop" / "Thesis" / "mednli" / "**" / "*.jsonl"), recursive=True):
        for l in open(f):
            r = json.loads(l); src.update(shingles(r.get("sentence1", "")) | shingles(r.get("sentence2", "")))
    out = set()              # shingles of the private model outputs
    priv = HERE / "results_private"
    for f in ("mimic_generations.jsonl", "mimic_qa.jsonl", "atomic_expert.jsonl", "atomic_main.jsonl"):
        if (priv / f).exists():
            for l in open(priv / f):
                out |= shingles(l)
    n_src, n_out = len(src), len(out - set(src))
    pub = HERE.parent / "mtsamples.csv"
    pub_sh = shingles(pub.read_text(errors="ignore")) if pub.exists() else set()
    tracked = subprocess.run(["git", "ls-files"], cwd=HERE, capture_output=True, text=True).stdout.split("\n")
    rows, all_src, all_out = [], set(), set()
    for t in [x for x in tracked if x]:
        sh = shingles(file_text(HERE / t)) - pub_sh
        m_src = [h for h in sh if h in src]
        all_src.update(m_src); all_out.update(h for h in sh if h in out and h not in src)
        rows.append(dict(file=t, shingles=len(sh), source_matches=len(m_src),
                         output_only_matches=sum(1 for h in sh if h in out and h not in src),
                         max_source_doc_freq=max((src[h] for h in m_src), default=0)))
    df = pd.DataFrame(rows).sort_values(["source_matches", "output_only_matches"], ascending=False)
    df.to_csv(HERE / "results" / "protected_text_scan.csv", index=False)
    hit_src = df[df.source_matches > 0].file.tolist()
    pilot = [f for f in hit_src if f.startswith("results/") and not f.startswith("results/mimic")]
    summary = [("tracked files scanned", len(df)), ("credentialed source shingles", n_src), ("model-output shingles", n_out),
               ("files with a source-text match", len(hit_src)), ("of which pilot result files", len(pilot)),
               ("distinct shingles shared with the source texts", len(all_src)),
               ("largest number of source documents containing one of them", max((src[h] for h in all_src), default=0)),
               ("files with a model-output match only", int(((df.source_matches == 0) & (df.output_only_matches > 0)).sum())),
               ("distinct shingles shared with the model outputs only", len(all_out))]
    if "--no-history" not in sys.argv:
        hrows = []
        for path, sha, text in history_texts():
            sh = shingles(text) - pub_sh
            m_src = [h for h in sh if h in src]
            hrows.append(dict(path=path, blob=sha[:10], shingles=len(sh), source_matches=len(m_src),
                              output_only_matches=sum(1 for h in sh if h in out and h not in src),
                              max_source_doc_freq=max((src[h] for h in m_src), default=0)))
        H = pd.DataFrame(hrows).sort_values(["source_matches", "output_only_matches"], ascending=False)
        H.to_csv(HERE / "results" / "protected_text_scan_history.csv", index=False)
        hs = H[H.source_matches > 0].path
        summary += [("file versions in the history scanned", len(H)), ("history versions with a source-text match", len(hs)),
                    ("of which versions of pilot result files", int(sum(1 for x in hs if x.startswith("results/") and not x.startswith("results/mimic")))),
                    ("largest source-document frequency in the history", int(H.max_source_doc_freq.max()) if len(H) else 0)]
        print(f"history: {len(H)} file versions scanned; with a source-text match: {int((H.source_matches > 0).sum())}; "
              f"largest source-document frequency of any match: {int(H.max_source_doc_freq.max()) if len(H) else 0}")
        print(H[(H.source_matches > 0) | (H.output_only_matches > 0)].head(25).to_string(index=False))
    pd.DataFrame(summary, columns=["check", "value"]).to_csv(HERE / "results" / "protected_text_scan_summary.csv", index=False)
    print(f"credentialed source shingles: {n_src:,}; additional model-output shingles: {n_out:,}; "
          f"public MTSamples shingles ignored: {len(pub_sh):,}")
    hit = df[(df.source_matches > 0) | (df.output_only_matches > 0)]
    print(f"tracked files scanned: {len(df)}; with a source-text match: {int((df.source_matches > 0).sum())}; "
          f"with a model-output match only: {int(((df.source_matches == 0) & (df.output_only_matches > 0)).sum())}")
    print(hit.to_string(index=False))


if __name__ == "__main__":
    main()
