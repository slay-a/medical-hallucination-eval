#!/usr/bin/env python3
"""
build_progress_slides.py — Fall 2026 progress report slides (PDF via reportlab, PPTX via python-pptx).
Numbers are read from ../results at build time.
"""
import subprocess
import sys
from datetime import date
from pathlib import Path

import pandas as pd
from reportlab.lib.colors import HexColor, black, white
from reportlab.lib.pagesizes import landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
RES = ROOT / "results"
sys.path.insert(0, str(ROOT / "thesis"))
from results_loader import Results, f3, f2, pct, pct0, fp  # noqa: E402

W, H = 13.333 * 72, 7.5 * 72
NAVY, ORANGE, BLUE, GREEN, GREY = HexColor("#1F2A44"), HexColor("#E07B54"), HexColor("#4C8BB5"), HexColor("#5DBF6E"), HexColor("#666666")
for name, fn in (("TNR", "Times New Roman.ttf"), ("TNRB", "Times New Roman Bold.ttf")):
    pdfmetrics.registerFont(TTFont(name, f"/System/Library/Fonts/Supplemental/{fn}"))


def slides(R: Results):
    from mimic_results import Mimic, NAME as MNAME, VAR as MVAR, JUDGE_NAMES
    M = Mimic()
    t10u, t10c, t20u, t20c = R.test("UFR", "E0", "E1"), R.test("CR", "E0", "E1"), R.test("UFR", "E0", "E2"), R.test("CR", "E0", "E2")
    cal = R.cal_row("all"); cov = R.cov.set_index("condition")
    var = R.cal_var[R.cal_var.group == "all"].set_index("variant") if R.cal_var is not None else None
    jc = pd.read_csv(RES / "judge_candidates.csv") if (RES / "judge_candidates.csv").exists() else None
    ftj = None
    ftp = Path.home() / "Desktop" / "Thesis" / "models" / "mednli-deberta-v3-large" / "mednli_finetune_summary.json"
    if ftp.exists():
        import json
        ftj = json.load(open(ftp))
    thesis_pdf = ROOT / "thesis" / "Ponangi_Thesis_Fall2026.pdf"
    pages = "—"
    try:
        out = subprocess.run(["pdfinfo", str(thesis_pdf)], capture_output=True, text=True).stdout
        pages = next(l.split(":")[1].strip() for l in out.splitlines() if l.startswith("Pages"))
    except Exception:  # noqa: BLE001
        pass
    best = None
    if jc is not None:
        best = jc[jc.group == "all"].sort_values("kappa", ascending=False).iloc[0]
    MODE = {"top3": "top-3 sentences", "doc": "whole course, windowed"}
    S = []
    S.append(dict(kind="title", title="Evaluating and Reducing Hallucinations in LLM-Based Medical Report Summarization",
                  sub=["Progress Report, Fall 2026 (COMP 698C)", "Srilaya Ponangi, M.S. Computer Science", "Advisor: Dr. Taehyung (George) Wang",
                       f"California State University, Northridge, {date.today():%B %d, %Y}"]))
    S.append(dict(kind="bullets", title="Goal and research questions (from the proposal)", bullets=[
        "Goal: measure hallucinations in LLM-generated patient-facing summaries at the claim level with a judge validated against medical experts, and test whether retrieval grounding and verification reduce them on real clinical notes",
        "RQ1: how often are LLM summary statements unsupported by, or contradicting, the source note?",
        "RQ2: which statements are unsupported, and how well does the automatic judge agree with medical experts?",
        "RQ3: do retrieval-augmented generation and verification reduce unsupported and contradicted statements, compared with an extractive bound, and at what cost in coverage?",
        "Two tasks in the proposal: summarization (main task) and document-grounded question answering (piloted with abstention)",
        "Three studies: pilot on public MTSamples notes with GPT-4o-mini; judge selection against expert labels; main study on MIMIC-IV hospital courses with a locally run open-weight model"]))
    judge_status = (f"Done: pilot judge validated (kappa {f2(cal.any_kappa)}); six candidates compared; {JUDGE_NAMES.get(best.judge, best.judge)} selected "
                    f"(kappa {f2(best.kappa)}, AUROC {f2(best.auroc)})" if best is not None else "Pilot judge validated; candidate comparison running")
    if ftj is not None:
        judge_status += f"; MedNLI fine-tune: dev accuracy {pct(ftj['dev_accuracy'])}"
    mimic_status = (f"Done: {M.n_docs} MIMIC-IV hospital courses, Qwen2.5-7B-Instruct served locally, no protected text leaves the laptop" if M.ok
                    else "Generation complete or in progress on 110 MIMIC-IV hospital courses with a local Qwen2.5-7B-Instruct; evaluation pending")
    S.append(dict(kind="table", title="Proposal commitments and where they stand", columns=["Commitment in the proposal", "Status"],
                  rows=[["Claim-level NLI judge calibrated against expert labels", judge_status],
                        ["Experiments on MIMIC data", mimic_status],
                        ["RAG, verification and an extractive baseline", "Done in both studies: E0 zero-context, E1 RAG (excerpts), E1b RAG (note + excerpts), E3 RAG + Chain-of-Verification, E2 extractive"],
                        ["Coverage against a reference", "Done: share of the clinician-written discharge-instruction sentences that the generated summary supports (judge run in reverse)"],
                        ["Citation accuracy", "Done: retrieval conditions cite an excerpt after every sentence; the judge checks each citation"],
                        ["Retrieval ablations (chunk size, top-k, BM25 vs dense, citations)", "Done: six E1 variants on the same documents"],
                        ["Question answering with abstention", "Pilot: three questions per course (medications, follow-up, warning signs), full course vs retrieved excerpts, 'Not stated in the note.' allowed"]],
                  widths=[4.6, 7.5]))
    S.append(dict(kind="image", title="Pipeline (identical judge for every condition)", image=RES / "fig_pipeline.png", caption="Every summary is scored by the same claim-level NLI judge; evidence comes from the same source document."))
    S.append(dict(kind="bullets", title="Data status", bullets=[
        f"MTSamples (public): {R.total_rows:,} transcriptions, {R.n_eligible} eligible consult and discharge notes, {R.n_docs} sampled (seed 42); {R.n_claims_total:,} pilot claims labeled",
        "ann-pt-summ v1.0.1 (PhysioNet, DUA signed): 210 expert-annotated patient summaries with 423 unsupported-fact spans for judge validation and selection; its 110 MIMIC-IV hospital courses with clinician-written discharge instructions are the main-study corpus",
        "MedNLI v1.0.0 (PhysioNet, DUA signed): 14,049 clinician-written premise and hypothesis pairs from MIMIC-III, used to adapt the judge",
        "MIMIC-IV-Note v2.2 (PhysioNet, credentialed): downloaded and verified in May 2026; the source of the hospital courses",
        "Ethics: CITI 'Data or Specimens Only Research' and 'Conflicts of Interest' completed April 10, 2026; every credentialed text is processed by models running on the author's laptop; nothing is sent to an API or to cloud storage"]))
    t1b0u, t30u, t30c, t31u = R.test("UFR", "E0", "E1b"), R.test("UFR", "E0", "E3"), R.test("CR", "E0", "E3"), R.test("UFR", "E1", "E3")
    conds = [c for c in ["E0", "E1", "E1b", "E3", "E2"] if c in R.conds]
    S.append(dict(kind="table", title=f"Pilot study: five conditions on {R.n_docs} MTSamples notes (GPT-4o-mini, MiniLM judge)",
                  columns=["Metric"] + [{"E0": "E0 zero-context", "E1": "E1 RAG excerpts", "E1b": "E1b RAG + note", "E3": "E3 RAG + CoVe", "E2": "E2 extractive"}[c] for c in conds],
                  rows=[["UFR mean"] + [f3(R.mean(c, "UFR")) for c in conds],
                        ["CR mean"] + [f3(R.mean(c, "CR")) for c in conds],
                        ["Supported share of claims"] + [pct0((R.label_counts.loc[c, "Supported"] / R.label_counts.loc[c].sum())) for c in conds],
                        ["Claims per summary"] + [f"{R.claims_per[c]:.1f}" for c in conds],
                        ["Words per summary"] + [f"{R.words[c].mean():.0f}" for c in conds],
                        ["Coverage of source sentences (cos ≥ 0.6)"] + [pct(cov.loc[c, "coverage_at_0_6"]) for c in conds]],
                  widths=[3.0] + [1.85] * len(conds),
                  note=(f"Paired Wilcoxon versus E0: E1 CR {fp(t10c.p)} (−{abs(t10c.rel_change_mean):.0f}%), E1 UFR {fp(t10u.p)}; "
                        f"E1b UFR {fp(t1b0u.p)}; E3 UFR {fp(t30u.p)} (−{abs(t30u.rel_change_mean):.0f}%; versus E1 {fp(t31u.p)}), E3 CR {fp(t30c.p)}; "
                        f"E2 UFR {fp(t20u.p)}. Verification lowers unsupported facts the most; retrieval alone lowers contradictions modestly; "
                        f"the extractive bound exposes a non-zero judge error floor. Header lines and abstentions excluded.")))
    S.append(dict(kind="table", title="Pilot headline: the small MiniLM judge disagrees with medical experts",
                  columns=["Group", "Sentences", "Expert-flagged", "Judge-flagged (UFR)", "Precision", "Recall", "Kappa", "AUROC"],
                  rows=[[r.group.replace("_", " "), f"{int(r.n_sentences):,}", pct(r.expert_flag_rate), pct(r.judge_UFR), f2(r.any_precision), f2(r.any_recall), f2(r.any_kappa), f2(r.auroc_1_minus_p_entail)]
                        for _, r in R.cal.iterrows() if r.group in ("all", "generated", "doctor_written", "gpt4_zero_shot", "llama_70b_original")],
                  widths=[2.2, 1.2, 1.5, 1.9, 1.2, 1.1, 1.0, 1.0],
                  note=f"On 1,781 expert-annotated sentences the pilot judge flags {pct0(cal.judge_UFR)} where experts flag {pct0(cal.expert_flag_rate)}; kappa {f2(cal.any_kappa)} is chance level; no aggregation variant raised kappa above {f2(var.best_kappa.max()) if var is not None else '—'}. This motivated the judge selection study."))
    if jc is not None:
        allr = jc[jc.group == "all"].sort_values("kappa", ascending=False)
        S.append(dict(kind="table", title="Judge selection: candidates scored against the expert annotations (1,781 sentences)",
                      columns=["Judge", "Evidence", "AUROC", "Precision", "Recall", "Kappa", "Flag rate"],
                      rows=[[JUDGE_NAMES.get(r.judge, r.judge), MODE.get(r["mode"], r["mode"]), f2(r.auroc), f2(r.precision), f2(r.recall), f2(r.kappa), pct0(r.flag_rate)] for _, r in allr.iterrows()],
                      widths=[4.0, 2.0, 1.1, 1.3, 1.1, 1.1, 1.3],
                      note=(f"Decision rule fixed in advance: highest kappa on all sentences, threshold chosen on the other subset. Selected: {JUDGE_NAMES.get(best.judge, best.judge)} "
                            f"({MODE.get(best['mode'])}), kappa {f2(best.kappa)}, AUROC {f2(best.auroc)}, precision {f2(best.precision)}, recall {f2(best.recall)}: moderate agreement, "
                            f"so main-study rates are estimates from an instrument with a known error profile, and the paired comparisons carry the weight.")))
        S.append(dict(kind="image", title="Judge selection: AUROC and kappa by candidate and evidence mode", image=RES / "fig_judges.png",
                      caption="Model size and training data matter more than the aggregation rule; whole-course evidence helps the large models."))
    S.append(dict(kind="bullets", title="Main study design: MIMIC-IV hospital courses", bullets=[
        "110 de-identified Brief Hospital Course sections (ann-pt-summ), each paired with the discharge instructions the treating clinician wrote",
        "Generator: Qwen2.5-7B-Instruct, 4-bit, served on the laptop with Apple MLX through an OpenAI-compatible interface; the pilot pipeline runs unchanged and no protected text leaves the machine",
        "Same five conditions; retrieval conditions must cite the excerpt behind every sentence; 'Not stated in the note.' allowed for unsupported sections",
        f"Judge: {JUDGE_NAMES.get(best.judge, best.judge) if best is not None else 'the selected candidate'} with whole-course evidence; UFR and CR per summary; the doctor-written instructions scored as a human reference",
        "Coverage measured against the clinician's own instructions (judge in reverse); citation accuracy; six retrieval ablations (3- or 8-sentence chunks, top-2 or top-5, BM25, no citations)",
        "Question-answering pilot: medications, follow-up and warning signs; full course versus three retrieved chunks; abstention rate, UFR/CR, and support by the clinician's instructions"]))
    if M.ok:
        mc = [c for c in ["E0", "E1", "E1b", "E3", "E2", "REF"] if c in M.conds]
        HEAD = {"E0": "E0 zero-context", "E1": "E1 RAG excerpts", "E1b": "E1b RAG + course", "E3": "E3 RAG + CoVe", "E2": "E2 extractive", "REF": "Clinician-written"}
        def cell(c, col, fmt):
            v = M.val(c, col)
            return "—" if v is None or (isinstance(v, float) and pd.isna(v)) else fmt(v)
        S.append(dict(kind="table", title=f"Main study results (n = {M.n_docs} hospital courses; judge: {M.judge_name()}, τ = {M.tau:.2f})",
                      columns=["Metric"] + [HEAD[c] for c in mc],
                      rows=[["UFR mean"] + [cell(c, "UFR_mean", f3) for c in mc],
                            ["CR mean"] + [cell(c, "CR_mean", f3) for c in mc],
                            ["Coverage of the clinician's sentences"] + [cell(c, "coverage_ref_mean", pct) if c != "REF" else "—" for c in mc],
                            ["Citation accuracy"] + [cell(c, "citation_accuracy_mean", pct) for c in mc],
                            ["Abstentions per summary"] + [cell(c, "abstentions_mean", lambda v: f"{v:.1f}") for c in mc],
                            ["Claims per summary"] + [cell(c, "claims_mean", lambda v: f"{v:.1f}") for c in mc],
                            ["Words per summary"] + [cell(c, "words_mean", lambda v: f"{v:.0f}") for c in mc]],
                      widths=[3.2] + [1.55] * len(mc),
                      note=" ".join(f"{c} vs E0: UFR {fp(M.test('UFR', c).p)}, CR {fp(M.test('CR', c).p)}, coverage {fp(M.test('coverage_ref', c).p)}." for c in mc if c not in ("E0", "REF") and M.test("UFR", c) is not None)
                           + " Paired two-sided Wilcoxon tests on the same documents."))
        S.append(dict(kind="image2", title="Main study: per-document distributions", images=[RES / "fig_mimic_box.png", RES / "fig_mimic_coverage.png"]))
        S.append(dict(kind="bullets", title="Main study: what the results say", bullets=[
            f"Every grounded condition cut the unsupported fact rate by about {M.rel(M.test('UFR','E1b'))} versus zero-context (E1 {f3(M.mean('E1','UFR'))}, E1b {f3(M.mean('E1b','UFR'))}, E3 {f3(M.mean('E3','UFR'))} versus E0 {f3(M.mean('E0','UFR'))}; all {fp(max(M.test('UFR', c).p for c in ('E1','E1b','E3')))}); the three do not differ from one another; CR did not change",
            f"Omission separates them: excerpt-only RAG covered {pct0(M.mean('E1','coverage_ref'))} of the clinician's sentences and verification {pct0(M.mean('E3','coverage_ref'))}, against {pct0(M.mean('E0','coverage_ref'))} for E0; RAG with the whole course kept {pct0(M.mean('E1b','coverage_ref'))} ({fp(M.test('coverage_ref','E1b').p)} versus E0)",
            f"Verbatim extraction scores UFR {f3(M.mean('E2','UFR'))}: the judge's error floor; the grounded LLM conditions sit within a few hundredths of it",
            f"The clinicians' own instructions score UFR {f3(M.mean('REF','UFR'))} under the same judge: support by the hospital course measures source-faithfulness, not clinical correctness",
            f"Citations: {pct0(M.mean('E1','citation_accuracy'))} accurate for excerpt-only RAG but {pct0(M.mean('E1b','citation_accuracy'))} for RAG with the whole course; {M.faithfulness_short()}; coverage followed the amount of retrieved text",
            (f"Question answering: abstention {pct0(M.qa_tot.loc['QA-E0','abstention_rate'])} (full course) versus {pct0(M.qa_tot.loc['QA-E1','abstention_rate'])} (excerpts), highest for warning signs, which the course rarely states; retrieval lowers unsupported answers partly by abstaining more" if M.qa_tot is not None else "")]))
        if M.abl is not None and len(M.abl):
            rows = []
            for v in [x for x in MVAR if x != "base"]:
                r = {m: M.ablation(m, v) for m in ("UFR", "CR", "coverage_ref", "citation_accuracy")}
                rows.append([MVAR[v]] + [(f"{r[m].delta_mean:+.3f} ({fp(r[m].p)})" if r[m] is not None and not pd.isna(r[m].p) else "—") for m in ("UFR", "CR", "coverage_ref", "citation_accuracy")])
            S.append(dict(kind="table", title="Retrieval ablations: paired difference from the E1 base configuration", columns=["Variant", "Δ UFR (p)", "Δ CR (p)", "Δ coverage (p)", "Δ citation accuracy (p)"],
                          rows=rows, widths=[3.6, 2.1, 2.1, 2.1, 2.2], note="Negative UFR and CR differences and positive coverage and citation differences favour the variant. Base: 5-sentence chunks, top-3, dense retrieval, citations required."))
        if M.qa_tot is not None:
            S.append(dict(kind="table", title="Question-answering pilot (three questions per course)", columns=["Condition", "Question", "n", "Abstained", "UFR", "CR", "Supported by clinician's instructions", "Words"],
                          rows=[[r.condition, r.question.replace("_", " "), str(int(r.n)), pct0(r.abstention_rate), f3(r.UFR_mean), f3(r.CR_mean), pct0(r.ref_support_mean), f"{r.words_mean:.0f}"]
                                for _, r in M.qa.sort_values(["condition", "question"]).iterrows()],
                          widths=[1.3, 1.6, 0.7, 1.2, 1.0, 1.0, 3.2, 0.9],
                          note="QA-E0 answers from the full hospital course; QA-E1 from the three chunks retrieved with the question. UFR and CR are computed on answered questions only."))
    else:
        S.append(dict(kind="bullets", title="Main study: status", bullets=[
            "Generation of the five conditions and six retrieval variants on the 110 hospital courses is running on the laptop (about five minutes per course with the local model)",
            "Evaluation with the selected judge, the question-answering pilot and the figures follow automatically; the thesis chapters read the result files when they exist"]))
    S.append(dict(kind="bullets", title="Thesis writing status", bullets=[
        f"Complete draft generated from the result files: {pages} pages in CSUN format (1-inch margins, 12-pt Times New Roman, double-spaced, roman-numbered preliminary pages, arabic body)",
        "Nine chapters: 1 Introduction, 2 Literature Review, 3 Data, 4 Methodology, 5 Pilot Study on MTSamples, 6 Selecting a Valid Judge, 7 Main Study on MIMIC-IV, 8 Discussion, 9 Conclusion; IEEE references with DOIs; Appendices A to E (prompts, per-document results, worked examples, repository, threshold sweep)",
        "Since the September 4 draft: judge selection study, MedNLI adaptation, main study on MIMIC-IV with a local model (citations, coverage against the clinician's instructions, citation accuracy, ablations), question-answering pilot",
        "Open items: committee composition to be confirmed with the advisor; advisor review of the draft; final title on the ETD planning form (with or without 'and Question Answering')",
        "Every table and figure is regenerated by one command from the result files"]))
    S.append(dict(kind="table", title="Plan to the defense (aligned with CSUN ETD deadlines)", columns=["By", "Milestone"],
                  rows=[["Sep 22, 2026", "Advisor review of the complete draft; committee composition discussed"],
                        ["Oct 9, 2026", "ETD Planning Form filed (title and committee)"],
                        ["Oct 16, 2026", "Revised draft to advisor for approval; Canvas upload and Turnitin check after approval"],
                        ["Oct 23, 2026", "Draft to committee (two-week review); Doodle poll with 1-hour slots over 3 to 4 weeks"],
                        ["Nov 6, 2026", "ETD preliminary format review"],
                        ["Nov 16 to 24, 2026", "Defense (40 min presentation, 10 min Q&A); slides to advisor one week before"],
                        ["Dec 4, 2026", "Final ETD upload after revisions"]], widths=[2.4, 9.7],
                  note="The December 12 defense date in the May report is after the December 4 ETD deadline and has been moved."))
    S.append(dict(kind="bullets", title="Risks and requests", bullets=[
        (f"Judge validity: the selected judge agrees with experts moderately (kappa {f2(best.kappa)}); main-study rates are reported with this error profile, and conclusions rest on paired comparisons" if best is not None
         else "Judge validity: the pilot judge agrees with experts at chance level; the judge selection study addresses this"),
        "Generator difference: the pilot used GPT-4o-mini on public notes and the main study a 7-billion-parameter local model on MIMIC-IV, so the two studies are compared qualitatively",
        "Schedule: advisor approval, Turnitin and two weeks of committee review must all precede the November 6 format review",
        "Requests to the advisor: review the complete draft, confirm the committee, and decide the final title for the ETD planning form"]))
    return S


# ───────────────────────────── PDF renderer ─────────────────────────────
def draw_chrome(c, i, n, title):
    c.setFillColor(white); c.rect(0, 0, W, H, fill=1, stroke=0)
    c.setFillColor(NAVY); c.rect(0, H - 62, W, 62, fill=1, stroke=0)
    c.setFillColor(white); c.setFont("TNRB", 22); c.drawString(36, H - 42, title)
    c.setFillColor(GREY); c.setFont("TNR", 9)
    c.drawString(36, 16, "Hallucination Evaluation in LLM-Based Medical Summarization · COMP 698C · Fall 2026")
    c.drawRightString(W - 36, 16, f"{i} / {n}")


def wrap(c, text, font, size, width):
    words = text.split(); lines, cur = [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if pdfmetrics.stringWidth(t, font, size) <= width:
            cur = t
        else:
            lines.append(cur); cur = w
    if cur:
        lines.append(cur)
    return lines


def draw_table(c, x, y, columns, rows, widths, font=9.5):
    scale = 72
    ws = [w * scale for w in widths]
    yy = y
    def row_height(vals, f):
        return max(len(wrap(c, str(v), "TNR", f, ws[j] - 8)) for j, v in enumerate(vals)) * (f + 3) + 6
    h = row_height(columns, font)
    c.setFillColor(HexColor("#DDE3EA")); c.rect(x, yy - h, sum(ws), h, fill=1, stroke=0)
    xx = x
    for j, v in enumerate(columns):
        c.setFillColor(black); c.setFont("TNRB", font)
        for k, line in enumerate(wrap(c, str(v), "TNRB", font, ws[j] - 8)):
            c.drawString(xx + 4, yy - 4 - (k + 1) * (font + 3) + 3, line)
        xx += ws[j]
    yy -= h
    for r in rows:
        h = row_height(r, font)
        c.setStrokeColor(HexColor("#BBBBBB")); c.line(x, yy, x + sum(ws), yy)
        xx = x
        for j, v in enumerate(r):
            c.setFillColor(black); c.setFont("TNR", font)
            for k, line in enumerate(wrap(c, str(v), "TNR", font, ws[j] - 8)):
                c.drawString(xx + 4, yy - 4 - (k + 1) * (font + 3) + 3, line)
            xx += ws[j]
        yy -= h
    c.line(x, yy, x + sum(ws), yy)
    return yy


def render_pdf(S, out: Path):
    c = canvas.Canvas(str(out), pagesize=(W, H))
    n = len(S)
    for i, s in enumerate(S, 1):
        if s["kind"] == "title":
            c.setFillColor(NAVY); c.rect(0, 0, W, H, fill=1, stroke=0); c.setFillColor(white)
            c.setFont("TNRB", 30); y = H - 200
            for line in wrap(c, s["title"], "TNRB", 30, W - 160):
                c.drawCentredString(W / 2, y, line); y -= 40
            c.setFont("TNR", 18); y -= 30
            for line in s["sub"]:
                c.drawCentredString(W / 2, y, line); y -= 28
            c.showPage(); continue
        draw_chrome(c, i, n, s["title"])
        if s["kind"] == "bullets":
            y = H - 105; c.setFillColor(black)
            for bl in s["bullets"]:
                c.setFont("TNR", 15); lines = wrap(c, bl, "TNR", 15, W - 130)
                c.setFillColor(ORANGE); c.circle(52, y + 5, 3.5, fill=1, stroke=0); c.setFillColor(black)
                for line in lines:
                    c.drawString(66, y, line); y -= 21
                y -= 10
        elif s["kind"] == "table":
            yy = draw_table(c, 36, H - 85, s["columns"], s["rows"], s["widths"], font=10.5)
            if s.get("note"):
                c.setFont("TNR", 11); c.setFillColor(NAVY); y = yy - 22
                for line in wrap(c, s["note"], "TNR", 11, W - 80):
                    c.drawString(36, y, line); y -= 15
        elif s["kind"] in ("image", "image2"):
            imgs = [s["image"]] if s["kind"] == "image" else s["images"]
            avail_w = (W - 72 - 20 * (len(imgs) - 1)) / len(imgs); avail_h = H - 62 - 80
            x = 36
            for im in imgs:
                if Path(im).exists():
                    ir = ImageReader(str(im)); iw, ih = ir.getSize(); sc = min(avail_w / iw, avail_h / ih)
                    c.drawImage(ir, x + (avail_w - iw * sc) / 2, H - 72 - ih * sc, iw * sc, ih * sc)
                x += avail_w + 20
            if s.get("caption"):
                c.setFont("TNR", 11); c.setFillColor(NAVY); y = 46
                for line in wrap(c, s["caption"], "TNR", 11, W - 80)[:3]:
                    c.drawString(36, y, line); y -= 14
        c.showPage()
    c.save()


def render_pptx(S, out: Path):
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for s in S:
        sl = prs.slides.add_slide(blank)
        tb = sl.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.9)); tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = s["title"]; p.runs[0].font.size = Pt(28 if s["kind"] != "title" else 30); p.runs[0].font.bold = True
        if s["kind"] == "title":
            for line in s["sub"]:
                q = tf.add_paragraph(); q.text = line; q.runs[0].font.size = Pt(18)
            continue
        if s["kind"] == "bullets":
            box = sl.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(12.3), Inches(5.8)); tf = box.text_frame; tf.word_wrap = True
            for i, bl in enumerate(s["bullets"]):
                q = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); q.text = "• " + bl; q.runs[0].font.size = Pt(16); q.space_after = Pt(8)
        elif s["kind"] == "table":
            rows, cols = len(s["rows"]) + 1, len(s["columns"])
            tbl = sl.shapes.add_table(rows, cols, Inches(0.4), Inches(1.2), Inches(12.5), Inches(0.4) * rows).table
            def set_cell(cell, text, size, bold=False):
                cell.text = str(text) if str(text).strip() else " "
                for run in cell.text_frame.paragraphs[0].runs:
                    run.font.size = Pt(size); run.font.bold = bold
            for j, cname in enumerate(s["columns"]):
                set_cell(tbl.cell(0, j), cname, 11, True)
            for i, r in enumerate(s["rows"], 1):
                for j, v in enumerate(r):
                    set_cell(tbl.cell(i, j), v, 10)
            total = sum(s["widths"])
            for j, w in enumerate(s["widths"]):
                tbl.columns[j].width = Inches(12.5 * w / total)
            if s.get("note"):
                nb = sl.shapes.add_textbox(Inches(0.4), Inches(6.5), Inches(12.5), Inches(0.8)); nb.text_frame.word_wrap = True
                nb.text_frame.paragraphs[0].text = s["note"]; nb.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        else:
            imgs = [s["image"]] if s["kind"] == "image" else s["images"]
            wpic = (12.5 - 0.3 * (len(imgs) - 1)) / len(imgs); x = 0.4
            for im in imgs:
                if Path(im).exists():
                    sl.shapes.add_picture(str(im), Inches(x), Inches(1.2), width=Inches(wpic))
                x += wpic + 0.3
            if s.get("caption"):
                nb = sl.shapes.add_textbox(Inches(0.4), Inches(6.6), Inches(12.5), Inches(0.7)); nb.text_frame.word_wrap = True
                nb.text_frame.paragraphs[0].text = s["caption"]; nb.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    prs.save(str(out))


def main():
    R = Results(); S = slides(R)
    render_pdf(S, HERE / "Progress_Slides_Fall2026.pdf"); render_pptx(S, HERE / "Progress_Slides_Fall2026.pptx")
    print(f"wrote {len(S)} slides: Progress_Slides_Fall2026.pdf / .pptx")


if __name__ == "__main__":
    main()
